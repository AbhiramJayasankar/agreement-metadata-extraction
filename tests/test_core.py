import json

from docx import Document
import pytest

from metadata_extractor.dataset import files_by_stem, read_labels, write_predictions
from metadata_extractor.evaluate import corrected_test_labels, recall_rows, usable_labels
from metadata_extractor.fields import CSV_COLUMNS, FACT_KEYS, clean_party, facts_to_fields, format_answer, infer_end_date
from metadata_extractor.llm import EXAMPLES, parse_json_object, select_examples
from metadata_extractor.text import extract_text


def test_docx_preserves_paragraph_and_table_order(tmp_path):
    document = Document()
    document.add_paragraph("  Before   the table ")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Tenant"
    document.add_paragraph("After")
    path = tmp_path / "agreement.docx"
    document.save(path)
    cache = tmp_path / "text"
    assert extract_text(path, cache) == "Before the table\nOwner\nTenant\nAfter"
    assert (cache / "agreement.txt").read_text(encoding="utf-8").endswith("After")


def test_compound_extension_matches_csv_stem(tmp_path):
    path = tmp_path / "example.pdf.docx"
    path.touch()
    assert files_by_stem(tmp_path) == {"example": path}


def test_json_in_fence_and_braces_in_name():
    values = dict.fromkeys(FACT_KEYS, "")
    values["lessor"] = "A {B}"
    assert parse_json_object("```json\n" + json.dumps(values) + "\n``` trailing")["lessor"] == "A {B}"


@pytest.mark.parametrize("output", ['{"lessor": "A"}', "no JSON", '{"broken"', '{"monthly_rent": []}'])
def test_invalid_json_or_schema_requires_retry(output):
    with pytest.raises(ValueError):
        parse_json_object(output)


@pytest.mark.parametrize("field,raw,expected", [
    ("Agreement Value", " Rs.9,000/- ", "9000"),
    ("Agreement Value", "6,500.00", "6500"),
    ("Agreement Value", "P6,500.50", "6500"),
    ("Renewal Notice (Days)", "60 days", "60"),
    ("Renewal Notice (Days)", "two months", ""),
    ("Agreement Start Date", "1/4/2010", "01.04.2010"),
    ("Agreement End Date", "31-2-2011", "31.02.2011"),
    ("Party One", "  Mr. A B  ", "Mr. A B"),
])
def test_formatting_does_not_infer_answers(field, raw, expected):
    assert format_answer(field, raw) == expected


def test_recall_empty_labels_and_lenient_names():
    gold = dict.fromkeys(CSV_COLUMNS, "") | {"File Name": "one", "Party One": " A.B "}
    prediction = dict.fromkeys(CSV_COLUMNS, "") | {"File Name": "one", "Party One": "a b"}
    scores, errors = recall_rows([gold], [prediction])
    assert scores[0]["Recall"] == 1
    assert scores[4]["Recall"] == 0
    assert scores[4]["Lenient Recall"] == 1
    assert len(errors) == 1
    missing_scores, _ = recall_rows([gold], [prediction], {"one"})
    assert all(row["Recall"] == 0 for row in missing_scores)


@pytest.mark.parametrize("start,months,expected", [
    ("01.04.2010", "11", "28.02.2011"),
    ("15.12.2012", "11", "14.11.2013"),
    ("31.01.2023", "1", "27.02.2023"),
    ("31.01.2024", "1", "28.02.2024"),
    ("29.02.2024", "12", "27.02.2025"),
    ("01.12.2023", "2", "31.01.2024"),
    ("", "11", ""), ("01.04.2010", "", ""),
    ("31.02.2011", "11", ""), ("01.04.2010", "0", ""),
])
def test_calendar_month_arithmetic(start, months, expected):
    assert infer_end_date(start, months) == expected


def test_facts_use_explicit_end_and_preserve_names():
    result = facts_to_fields({
        "monthly_rent": "6,500.00", "start_date": "1/4/2010", "end_date": "30/3/2011",
        "term_months": "11", "notice_amount": "2", "notice_unit": "months",
        "lessor": "  A.B  & C.D  ", "lessee": "UPPER Name",
    })
    assert list(result) == ["Agreement Value", "Agreement Start Date", "Agreement End Date", "Renewal Notice (Days)", "Party One", "Party Two"]
    assert result["Agreement Value"] == "6500"
    assert result["Agreement End Date"] == "30.03.2011"
    assert result["Renewal Notice (Days)"] == "60"
    assert result["Party One"] == "A.B  & C.D"
    assert result["Party Two"] == "UPPER Name"


@pytest.mark.parametrize("amount,unit,expected", [
    ("15", "days", "15"), ("3", "months", "90"), ("0", "days", "0"),
    ("1", "month", "30"), ("1", "Day", "1"),
    ("", "months", ""), ("2", "", ""), ("two", "months", ""),
])
def test_notice_arithmetic(amount, unit, expected):
    assert facts_to_fields({"notice_amount": amount, "notice_unit": unit})["Renewal Notice (Days)"] == expected


def test_missing_facts_remain_empty():
    assert all(value == "" for value in facts_to_fields({}).values())


def test_examples_are_leave_one_out():
    for example in EXAMPLES:
        stem = example["file"].removesuffix(".docx").removesuffix(".png")
        selected = select_examples(stem)
        assert len(selected) == 2
        assert example not in selected
        assert all(set(item["facts"]) == set(FACT_KEYS) for item in selected)


def test_usable_view_excludes_only_known_alignment_errors():
    names = ["24158401-Rental-Agreement", "54770958-Rental-Agreement", "54945838-Rental-Agreement", "18325926-Rental-Agreement-1"]
    assert usable_labels([{"File Name": name} for name in names]) == [{"File Name": names[-1]}]


def test_gold_dates_are_not_validated_or_repaired(tmp_path):
    path = tmp_path / "labels.csv"
    row = dict.fromkeys(CSV_COLUMNS, "") | {"File Name": "one", "Aggrement End Date": " 31.02.2011 "}
    write_predictions(path, [row])
    assert read_labels(path)[0]["Aggrement End Date"] == "31.02.2011"


@pytest.mark.parametrize("name,expected", [
    ("SRI VYSHNAVI DAIRY SPECIALITIES Private Ltd.", "VYSHNAVI DAIRY SPECIALITIES Private Ltd"),
    ("Mr. P. JohnsonRavikumar", "P. JohnsonRavikumar"),
    ("Mrs.Asha Ramesh", "Asha Ramesh"), ("B.Kishore", "B.Kishore"),
    ("Drona Rao", "Drona Rao"), ("Mrinal Sen", "Mrinal Sen"),
    ("K. Parthasarathy", "K. Parthasarathy"), ("", ""),
    (" shRI. A.B,;: ", "A.B"), ("Dr Mr. Name", "Mr. Name"),
])
def test_clean_party(name, expected):
    assert clean_party(name) == expected
    assert facts_to_fields({"lessor": name, "lessee": name})["Party One"] == expected
    assert facts_to_fields({"lessor": name, "lessee": name})["Party Two"] == expected


def test_split_pages_cuts_on_white_row():
    from PIL import Image, ImageDraw
    from metadata_extractor.ocr import split_pages

    image = Image.new("RGB", (200, 532), "white")
    draw = ImageDraw.Draw(image)
    for top in (30, 100, 186, 300, 400):
        draw.rectangle((10, top, 190, top + 15), fill="black")
    pages = split_pages(image)
    assert len(pages) == 2
    cut = pages[0].height
    assert 186 <= cut < 346
    assert all(image.getpixel((x, cut)) == (255, 255, 255) for x in range(image.width))
    assert sum(page.height for page in pages) == image.height


def test_image_cache_separates_engines(tmp_path, monkeypatch):
    from PIL import Image
    from metadata_extractor import text

    path = tmp_path / "scan.png"
    Image.new("RGB", (10, 10), "white").save(path)
    cache = tmp_path / "text"
    cache.mkdir()
    (cache / "scan.txt").write_text("old OCR", encoding="utf-8")
    calls = []
    def fake_ocr(path, engine):
        calls.append(engine)
        return engine
    monkeypatch.setattr(text, "image_text", fake_ocr)
    for engine in ("glm", "tesseract", "glm", "tesseract"):
        assert extract_text(path, cache, engine=engine) == engine
        assert (cache / f"scan.{engine}.txt").read_text(encoding="utf-8") == engine
    assert calls == ["glm", "tesseract"]


def test_corrected_test_scoring_does_not_change_original_labels():
    labels = [dict.fromkeys(CSV_COLUMNS, "") | {
        "File Name": "228094620-Rental-Agreement", "Party Two": ".B.Kishore",
    }, dict.fromkeys(CSV_COLUMNS, "") | {
        "File Name": "95980236-Rental-Agreement", "Aggrement End Date": "31.03.2011",
    }]
    predictions = [dict(labels[0], **{"Party Two": "B.Kishore"}),
                   dict(labels[1], **{"Aggrement End Date": "28.02.2011"})]
    scores, errors = recall_rows(labels, predictions)
    assert len(errors) == 2
    assert scores[5]["Recall"] == 0.5
    assert scores[5]["Lenient Recall"] == 1
    corrected_scores, corrected_errors = recall_rows(corrected_test_labels(labels), predictions)
    assert corrected_errors == []
    assert all(row["Recall"] == 1 for row in corrected_scores)
    assert labels[0]["Party Two"] == ".B.Kishore"
    assert labels[1]["Aggrement End Date"] == "31.03.2011"
